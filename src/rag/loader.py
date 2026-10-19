# ============================================================
# 文档加载 + 切块
# ------------------------------------------------------------
# RAG 第一步：把原始文档读进来，切成适合检索的小块（chunk）。
#
# 【本版重点修复】"知识库里明明有答案，却只检索到一部分"
#   旧版对所有文件一律定长硬切（400 字一刀），带来两类问题：
#     1) 小文件被切碎 —— 简历只有 1372 字却被切成 5 块，
#        "两段实习经历"正好被劈在相邻两块里，只召回其中一块就丢信息；
#     2) 句子被拦腰截断 —— 硬切不看标点，一句话被切成两半，
#        单块语义不完整，检索和生成都会失准。
#   对策：
#     a) 小文件（默认 ≤2000 字）整篇作为一个块，根本不切；
#     b) 大文件按"段落 → 句子"的语义边界聚合切块，尽量不切断完整句子；
#     c) PDF 提取出的中文常被插入空格（如"摘 要"），统一归一化，
#        否则用户问"摘要"永远匹配不到。
#
# 【重要】本模块支持的格式必须与 web_server.py 的上传白名单保持一致，
# 否则会出现"文件上传成功却检索不到"的静默 bug（历史教训）。
# 支持的格式：.txt .md .markdown .csv .json .pdf .docx
# ============================================================
from __future__ import annotations

import csv
import io
import json
import re
from dataclasses import dataclass
from pathlib import Path

# 支持的文件类型（与上传白名单统一维护）
TEXT_EXTS = {".txt", ".md", ".markdown"}
CSV_EXTS = {".csv"}
JSON_EXTS = {".json"}
PDF_EXTS = {".pdf"}
DOCX_EXTS = {".docx"}
SUPPORTED_EXTS = TEXT_EXTS | CSV_EXTS | JSON_EXTS | PDF_EXTS | DOCX_EXTS

# 小于该字数的文件整篇作为一个块，不切。
# 为什么：简历 / FAQ / 政策这类小文件整篇也就一两千字，塞进上下文成本很低；
# 但一旦被切碎，跨块的信息就会检索不全——这正是"两段实习只找到一段"的根因。
DEFAULT_MIN_CHUNK_FILE = 2000

# 中文（含中文标点）之间被 PDF 排版插入的空格
_CJK_SPACE_RE = re.compile(
    r"(?<=[\u4e00-\u9fff\u3000-\u303f\uff00-\uffef])[ \t]+"
    r"(?=[\u4e00-\u9fff\u3000-\u303f\uff00-\uffef])"
)
# 句末标点（用于句子边界切分）
_SENTENCE_END = "。！？；!?;"


@dataclass
class Document:
    """一个文档块：内容 + 来源元数据（企业场景中溯源非常重要）。"""
    content: str
    source: str          # 来源文件
    chunk_id: int        # 在原文中的第几块（同源块按 chunk_id 连续，可用于邻块扩展）


# ------------------------------------------------------------
# 文本归一化
# ------------------------------------------------------------
def normalize_text(text: str) -> str:
    """统一换行、去掉中文之间被 PDF 插入的多余空格、压缩连续空行。

    例：'摘 要' -> '摘要'，'实 习 经 历' -> '实习经历'
    英文单词之间的空格不受影响（只处理"中文—空格—中文"的情形）。
    """
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = _CJK_SPACE_RE.sub("", text)
    text = re.sub(r"[ \t]+\n", "\n", text)      # 去行尾空白
    text = re.sub(r"\n{3,}", "\n\n", text)      # 连续空行压成一个
    return text.strip()


# ------------------------------------------------------------
# 各格式的解析函数：统一返回"纯文本"
# ------------------------------------------------------------
def _read_text(path: Path) -> str:
    """读取纯文本文件（自动兼容 UTF-8 / GBK）。"""
    try:
        return path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        # 兼容 Windows 下常见的 GBK 编码文件
        return path.read_text(encoding="gbk", errors="ignore")


def _read_pdf(path: Path) -> str:
    """解析 PDF，抽取所有页的文本（保留页码分隔，便于人工核对）。"""
    try:
        from pypdf import PdfReader
    except ImportError:
        return ""   # 未安装解析库时返回空（上层会提示）
    reader = PdfReader(str(path))
    pages = []
    for i, page in enumerate(reader.pages, 1):
        pages.append(f"\n=== 第 {i} 页 ===\n" + (page.extract_text() or ""))
    return "\n".join(pages)


def _read_docx(path: Path) -> str:
    """解析 Word 文档，抽取所有段落的文本。"""
    try:
        import docx
    except ImportError:
        return ""
    doc = docx.Document(str(path))
    return "\n".join(p.text for p in doc.paragraphs)


def _read_csv(path: Path) -> str:
    """解析 CSV，转成"列名: 值"的可读文本（比原始逗号分隔更利于检索）。"""
    text = _read_text(path)
    reader = csv.reader(io.StringIO(text))
    rows = list(reader)
    if not rows:
        return ""
    header = rows[0]
    lines = []
    for row in rows[1:]:
        # 把每行拼成 "字段1: 值1 | 字段2: 值2" 的形式
        pairs = [f"{h}: {v}" for h, v in zip(header, row)]
        lines.append(" | ".join(pairs))
    return "\n".join(lines)


def _read_json(path: Path) -> str:
    """解析 JSON，扁平化成可读文本。"""
    try:
        data = json.loads(_read_text(path))
    except Exception:
        return _read_text(path)   # 解析失败就按纯文本处理
    return json.dumps(data, ensure_ascii=False, indent=2)


# 格式 -> 解析函数 的映射表（新增格式只需在这里加一行）
_PARSERS = {
    **{ext: _read_text for ext in TEXT_EXTS},
    **{ext: _read_csv for ext in CSV_EXTS},
    **{ext: _read_json for ext in JSON_EXTS},
    **{ext: _read_pdf for ext in PDF_EXTS},
    **{ext: _read_docx for ext in DOCX_EXTS},
}


def load_text_files(data_dir: str) -> list[tuple[str, str]]:
    """加载目录下所有受支持的文档，统一转成纯文本（已归一化）。

    返回: [(文件名, 全文内容), ...]
    """
    docs: list[tuple[str, str]] = []
    root = Path(data_dir)
    if not root.exists():
        return docs
    for path in sorted(root.rglob("*")):
        if not path.is_file():
            continue
        ext = path.suffix.lower()
        parser = _PARSERS.get(ext)
        if parser is None:
            continue   # 不支持的格式直接跳过
        try:
            text = parser(path)
        except Exception as e:
            print(f"[loader] 解析 {path.name} 失败：{e}")
            continue
        if text and text.strip():
            docs.append((path.name, normalize_text(text)))
    return docs


# ------------------------------------------------------------
# 切块
# ------------------------------------------------------------
def split_text(text: str, chunk_size: int = 500, chunk_overlap: int = 50) -> list[str]:
    """按字符长度硬切（定长切块）。

    保留这个最朴素的实现有两个用途：
      1) 单元测试直接验证切块数学逻辑；
      2) 语义切块里遇到"单句就超长"时的兜底。
    正式建索引请用 split_text_smart()。
    """
    if chunk_size <= 0:
        raise ValueError("chunk_size 必须为正数")
    if chunk_overlap >= chunk_size:
        raise ValueError("chunk_overlap 必须小于 chunk_size，否则会死循环")

    chunks: list[str] = []
    start = 0
    step = chunk_size - chunk_overlap  # 每次前进的步长
    while start < len(text):
        chunks.append(text[start:start + chunk_size])
        start += step
    return chunks


def _semantic_units(text: str, chunk_size: int) -> list[str]:
    """把长文本拆成一串"语义单元"，每个单元长度 ≤ chunk_size。

    优先按空行分段；段落仍超长时按句末标点断句再聚合。
    这样切块边界基本落在段落/句子之间，而不是随机切在某个字上。
    """
    units: list[str] = []
    for para in re.split(r"\n\s*\n", text):
        para = para.strip()
        if not para:
            continue
        if len(para) <= chunk_size:
            units.append(para)
            continue

        # 段落超长：按句末标点断句（保留标点，用零宽断言）
        buf = ""
        for sent in re.split(rf"(?<=[{_SENTENCE_END}])", para):
            if not sent:
                continue
            if len(sent) > chunk_size:
                # 单句就超长（PDF 提取常见：整页没标点），只能硬切兜底
                if buf:
                    units.append(buf)
                    buf = ""
                units.extend(split_text(sent, chunk_size, 0))
                continue
            if len(buf) + len(sent) <= chunk_size:
                buf += sent
            else:
                if buf:
                    units.append(buf)
                buf = sent
        if buf:
            units.append(buf)
    return units


def split_text_smart(text: str, chunk_size: int = 800, chunk_overlap: int = 120) -> list[str]:
    """语义切块：先按段落/句子切成语义单元，再带 overlap 聚合成块。

    相比硬切的收益：
      - 块内是一段（或几段）完整的话，"一件事"不会被劈成两半；
      - 块数更少 => 建索引的嵌入请求更少 => 省 token；
      - 检索命中后语义更完整，模型不容易答"上下文不支持"。
    """
    if chunk_size <= 0:
        raise ValueError("chunk_size 必须为正数")
    if chunk_overlap >= chunk_size:
        raise ValueError("chunk_overlap 必须小于 chunk_size，否则会死循环")

    units = _semantic_units(text, chunk_size)
    if not units:
        return []

    chunks: list[str] = []
    buf = ""
    for u in units:
        if not buf:
            buf = u
        elif len(buf) + 1 + len(u) <= chunk_size:
            buf = f"{buf}\n{u}"
        else:
            chunks.append(buf)
            # overlap：把上一块的尾部接过来，避免跨块信息断档
            tail = buf[-chunk_overlap:] if chunk_overlap > 0 else ""
            buf = f"{tail}\n{u}" if tail else u
            if len(buf) > chunk_size:
                # 极少数情况：overlap 尾巴 + 单元超长，退回硬切
                hard = split_text(buf, chunk_size, chunk_overlap)
                chunks.extend(hard[:-1])
                buf = hard[-1]
    if buf:
        chunks.append(buf)
    return chunks


def load_and_split(
    data_dir: str,
    chunk_size: int,
    chunk_overlap: int,
    min_file_chars: int = DEFAULT_MIN_CHUNK_FILE,
) -> list[Document]:
    """组合函数：把目录里的文档读入并切成 Document 列表。

    参数:
        data_dir      : 知识库目录
        chunk_size    : 大文件的块大小（字符）
        chunk_overlap : 块间重叠字符数
        min_file_chars: 文件短于该字数则整篇作为一个块（不切）
    """
    documents: list[Document] = []
    for filename, text in load_text_files(data_dir):
        text = text.strip()
        if not text:
            continue
        # 小文件整篇一块：避免"两段实习经历被劈成两块只召回一块"这类问题
        if min_file_chars > 0 and len(text) <= min_file_chars:
            documents.append(Document(content=text, source=filename, chunk_id=0))
            continue
        for i, chunk in enumerate(split_text_smart(text, chunk_size, chunk_overlap)):
            chunk = chunk.strip()
            if chunk:  # 丢掉空白块
                documents.append(Document(content=chunk, source=filename, chunk_id=i))
    return documents
